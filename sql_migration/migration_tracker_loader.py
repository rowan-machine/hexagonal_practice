"""
SQL Migration Tracker Loader and Manager

Loads and manages the CSV-based SQL migration tracker.
Provides utilities for tracking migration progress across 100+ SQL scripts.

Usage:
    from sql_migration.migration_tracker_loader import MigrationTracker
    
    tracker = MigrationTracker()
    tracker.load()
    tracker.get_pending_queries()
    tracker.update_status('CLAIMS_MEMBER_001', 'In Progress')
"""
import csv
import os
from typing import List, Dict, Any, Optional
from pathlib import Path
from datetime import datetime


class MigrationTracker:
    """
    Manages SQL migration tracking from CSV file.
    
    Provides methods to:
    - Load migration tracker CSV
    - Query by status, priority, etc.
    - Update migration status
    - Generate progress reports
    """
    
    def __init__(self, tracker_path: str = "sql_migration/migration_tracker.csv"):
        """
        Initialize migration tracker.
        
        Args:
            tracker_path: Path to CSV tracker file
        """
        self.tracker_path = Path(tracker_path)
        self.queries: List[Dict[str, Any]] = []
    
    def load(self) -> List[Dict[str, Any]]:
        """
        Load migration tracker from CSV.
        
        Returns:
            List of query dictionaries
        """
        if not self.tracker_path.exists():
            raise FileNotFoundError(f"Tracker file not found: {self.tracker_path}")
        
        self.queries = []
        with open(self.tracker_path, 'r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                self.queries.append(row)
        
        return self.queries
    
    def save(self) -> None:
        """Save migration tracker to CSV."""
        if not self.queries:
            return
        
        # Get fieldnames from first query
        fieldnames = list(self.queries[0].keys())
        
        with open(self.tracker_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(self.queries)
    
    def get_by_status(self, status: str) -> List[Dict[str, Any]]:
        """Get queries by status (Complete, Pending, In Progress, etc.)."""
        return [q for q in self.queries if q.get('status', '').lower() == status.lower()]
    
    def get_by_priority(self, priority: str) -> List[Dict[str, Any]]:
        """Get queries by priority (High, Medium, Low)."""
        return [q for q in self.queries if q.get('priority', '').lower() == priority.lower()]
    
    def get_pending_queries(self) -> List[Dict[str, Any]]:
        """Get all pending queries."""
        return self.get_by_status('Pending')
    
    def get_complete_queries(self) -> List[Dict[str, Any]]:
        """Get all completed queries."""
        return self.get_by_status('Complete')
    
    def update_status(self, query_id: str, status: str, 
                     assigned_to: Optional[str] = None,
                     notes: Optional[str] = None) -> bool:
        """
        Update query status.
        
        Args:
            query_id: Query ID to update
            status: New status (Complete, Pending, In Progress, etc.)
            assigned_to: Person assigned (optional)
            notes: Additional notes (optional)
        
        Returns:
            True if updated, False if query_id not found
        """
        for query in self.queries:
            if query.get('query_id') == query_id:
                query['status'] = status
                if assigned_to:
                    query['assigned_to'] = assigned_to
                if notes:
                    query['notes'] = notes
                if status == 'Complete':
                    query['completed_date'] = datetime.now().strftime('%Y-%m-%d')
                self.save()
                return True
        return False
    
    def get_progress_summary(self) -> Dict[str, Any]:
        """Get migration progress summary."""
        total = len(self.queries)
        complete = len(self.get_complete_queries())
        pending = len(self.get_pending_queries())
        in_progress = len(self.get_by_status('In Progress'))
        
        return {
            'total': total,
            'complete': complete,
            'pending': pending,
            'in_progress': in_progress,
            'percent_complete': (complete / total * 100) if total > 0 else 0
        }
    
    def add_query(self, query_data: Dict[str, Any]) -> None:
        """Add a new query to the tracker."""
        # Ensure required fields
        required = ['query_id', 'query_name', 'status', 'priority']
        for field in required:
            if field not in query_data:
                raise ValueError(f"Missing required field: {field}")
        
        # Set defaults
        query_data.setdefault('original_sql_file', '')
        query_data.setdefault('business_rule_location', '')
        query_data.setdefault('pipeline_usage', '')
        query_data.setdefault('sdk_usage', '')
        query_data.setdefault('test_coverage', 'No')
        query_data.setdefault('notes', '')
        query_data.setdefault('estimated_effort_days', '')
        query_data.setdefault('actual_effort_days', '')
        query_data.setdefault('assigned_to', '')
        query_data.setdefault('completed_date', '')
        query_data.setdefault('validation_status', '')
        
        self.queries.append(query_data)
        self.save()
    
    def export_to_excel(self, output_path: str = "sql_migration/migration_tracker.xlsx") -> None:
        """
        Export tracker to Excel format.
        
        Requires: pip install openpyxl
        """
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Migration Tracker"
            
            # Headers
            if self.queries:
                headers = list(self.queries[0].keys())
                ws.append(headers)
                
                # Style headers
                header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                
                # Data rows
                for query in self.queries:
                    row = [query.get(h, '') for h in headers]
                    ws.append(row)
                    
                    # Color code by status
                    status = query.get('status', '').lower()
                    if status == 'complete':
                        ws.cell(row=ws.max_row, column=1).fill = PatternFill(
                            start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"
                        )
                    elif status == 'in progress':
                        ws.cell(row=ws.max_row, column=1).fill = PatternFill(
                            start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"
                        )
            
            wb.save(output_path)
            print(f"✅ Exported to Excel: {output_path}")
        except ImportError:
            print("⚠️  openpyxl not installed. Install with: pip install openpyxl")
            print(f"   CSV file available at: {self.tracker_path}")


if __name__ == "__main__":
    # Example usage
    tracker = MigrationTracker()
    tracker.load()
    
    summary = tracker.get_progress_summary()
    print("=" * 60)
    print("SQL Migration Progress")
    print("=" * 60)
    print(f"Total Queries: {summary['total']}")
    print(f"Complete: {summary['complete']}")
    print(f"In Progress: {summary['in_progress']}")
    print(f"Pending: {summary['pending']}")
    print(f"Progress: {summary['percent_complete']:.1f}%")
    print("=" * 60)
    
    # Export to Excel if available
    tracker.export_to_excel()

